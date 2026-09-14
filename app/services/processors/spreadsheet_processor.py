from pathlib import Path

from openpyxl import load_workbook

from app.services.processors.base import FileProcessor


class SpreadsheetProcessor(FileProcessor):
    
    SPREADSHEET_TYPES = {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    }
    
    def supports(self, mime_type: str,) -> bool:
        return mime_type in self.SPREADSHEET_TYPES
    
    def process(self, path: Path, filename: str, mime_type: str):
        
        workbook = load_workbook(filename=path, read_only=True, data_only=False)
        
        sheets = []
        
        for worksheet in workbook.worksheets:
            rows = []
            
            for row in worksheet.iter_rows(values_only=True):
                values = [self.normalize_value(value) for value in row]
                
                if any(value is not None for value in values):
                    rows.append(values)

            sheets.append({
                "sheet_name": worksheet.title,
                "rows": rows,
            })